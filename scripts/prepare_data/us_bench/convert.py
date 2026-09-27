"""Convert the public ultrasound segmentation / detection benchmarks to the common layout
(`common.py`). Raw downloads: `download.sh` in the raw folder (sources verified 2026-09-26).

    python scripts/prepare_data/us_bench/convert.py --raw $BIOMLLM_DATA/raw/us_bench \
        --dst $BIOMLLM_DATA/us_bench --datasets all

| name           | organ / target                    | split                                   | group     |
|----------------|-----------------------------------|-----------------------------------------|-----------|
| busbra         | breast lesion                     | official 5-fold CV: fold1 test, fold2 val | patient   |
| busi           | breast lesion (+ normal images)   | seeded 70/10/20                         | image*    |
| bus_uclm       | breast lesion (+ normal images)   | seeded 70/10/20                         | patient   |
| breast_lesions | breast lesion (BrEaST, TCIA)      | seeded 70/10/20                         | case      |
| tn3k           | thyroid nodule                    | official test; val = 10% of trainval     | image*    |
| mmotu          | ovarian tumour                    | official val -> test; val = 10% of train | image*    |
| hc18           | fetal head                        | seeded 70/10/20 (labelled train only)    | fetus     |
| psfhs          | pubic symphysis + fetal head      | seeded 70/10/20                          | image*    |
| camus          | LV endocardium, myocardium, LA    | seeded 70/10/20                          | patient   |

(*) the source has no patient id: near-duplicates across splits are measured by the dHash audit
in meta.json instead.
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import re
from pathlib import Path

import numpy as np
from PIL import Image

from common import Writer, fill_ellipse_contour

LESION = True      # one box per connected component
STRUCTURE = False  # one box per structure


def _gray_rgb(a: np.ndarray) -> Image.Image:
    a = np.asarray(a)
    if a.dtype != np.uint8:
        a = a.astype(np.float32)
        a = (255 * (a - a.min()) / max(1e-6, a.max() - a.min())).astype(np.uint8)
    return Image.fromarray(a).convert("RGB")


def busbra(raw: Path, dst: Path) -> dict:
    root = raw / "busbra/x/BUSBRA"
    info = {r["ID"]: r for r in csv.DictReader(open(root / "bus_data.csv"))}
    fold = {r["ID"]: int(r["kFold"]) for r in csv.DictReader(open(root / "5-fold-cv.csv"))}
    w = Writer(dst, ["breast lesion"], LESION)
    for sid, r in sorted(info.items()):
        m = np.array(Image.open(root / "Masks" / f"mask_{sid[4:]}.png")) > 0
        split = {1: "test", 2: "val"}.get(fold[sid], "train")
        w.add(sid, Image.open(root / "Images" / f"{sid}.png"), m.astype(np.uint8), group=f"case{r['Case']}",
              split=split, extra={"pathology": r["Pathology"], "birads": r["BIRADS"], "device": r["Device"]})
    return w.finish({"source": "https://zenodo.org/records/8231412", "license": "CC BY 4.0",
                     "citation": "Gomez-Flores et al., BUS-BRA, Medical Physics 2024",
                     "split": "official 5-fold patient-level CV (5-fold-cv.csv, kFold): fold 1 = test, fold 2 = val"})


def busi(raw: Path, dst: Path) -> dict:
    import zipfile

    x = raw / "busi/x"
    if not x.exists():
        zipfile.ZipFile(raw / "busi/busi.zip").extractall(x)
    root = next(x.rglob("Dataset_BUSI_with_GT"))
    w = Writer(dst, ["breast lesion"], LESION)
    for cls in ("benign", "malignant", "normal"):
        for f in sorted((root / cls).glob("*.png")):
            if "_mask" in f.name:
                continue
            masks = sorted((root / cls).glob(f"{f.stem}_mask*.png"))
            img = Image.open(f)
            m = np.zeros((img.height, img.width), np.uint8)
            for mf in masks:
                a = np.array(Image.open(mf).convert("L")) > 127
                m |= a.astype(np.uint8)
            w.add(f"{cls}_{f.stem}", img, m, group=f"{cls}_{f.stem}", extra={"pathology": cls,
                                                                            "n_mask_files": len(masks)})
    return w.finish({"source": "HF gymprathap/Breast-Cancer-Ultrasound-Images-Dataset (mirror of Kaggle BUSI; "
                               "file list identical to the original)", "license": "CC BY 4.0 (per mirror card)",
                     "citation": "Al-Dhabyani et al., Data in Brief 2020",
                     "split": "seeded (0) 70/10/20 by near-duplicate cluster (dHash <= 6); no patient ids in the source",
                     "notes": "multiple *_mask_k.png files merged (OR); normal images kept with empty masks; "
                              "BUSI is known to contain near-duplicates and burnt-in annotations"}, dedup=True)


def bus_uclm(raw: Path, dst: Path) -> dict:
    root = next((raw / "bus_uclm/x").rglob("INFO.csv")).parent
    rows = list(csv.DictReader(open(root / "INFO.csv", encoding="utf-8-sig"), delimiter=";"))
    w = Writer(dst, ["breast lesion"], LESION)
    dropped = []
    for r in rows:
        name = r["Image"]
        if Image.open(root / "images" / name).size != Image.open(root / "masks" / name).size:
            dropped.append(name)  # images cropped w.r.t. their masks: alignment unknown
            continue
        m = np.array(Image.open(root / "masks" / name).convert("RGB"))
        lesion = (m.max(-1) > 127).astype(np.uint8)
        w.add(Path(name).stem, Image.open(root / "images" / name), lesion, group=name.split("_")[0],
              extra={"pathology": r["Label"].lower(), "doppler": r["Doppler"], "marks": r["Marks"]})
    return w.finish({"source": "https://data.mendeley.com/datasets/7fvgj4jsp7", "license": "CC BY 4.0",
                     "citation": "Vallez et al., BUS-UCLM, Scientific Data 2025",
                     "split": "seeded (0) 70/10/20 by patient (filename prefix)",
                     "notes": "RGB masks (green benign, red malignant) merged into one lesion class; "
                              "pathology kept per record; normal images kept with empty masks; "
                              f"{len(dropped)} images dropped (image cropped w.r.t. mask, alignment unknown): "
                              f"{sorted(dropped)}"})


def breast_lesions(raw: Path, dst: Path) -> dict:
    root = raw / "breast_lesions/x/BrEaST-Lesions_USG-images_and_masks"
    clin = {}
    try:
        import openpyxl

        ws = openpyxl.load_workbook(raw / "breast_lesions/clinical.xlsx", read_only=True).active
        head = [str(c.value) for c in next(ws.iter_rows(max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            d = dict(zip(head, row))
            case = str(d.get("Image_filename") or d.get("CaseID") or "")
            if case:
                clin[Path(case).stem] = {k: str(v) for k, v in d.items()
                                         if k in ("Classification", "Diagnosis", "BIRADS", "Pathology")}
    except Exception as e:  # clinical metadata is optional
        clin = {"_error": str(e)}
    w = Writer(dst, ["breast lesion"], LESION)
    for f in sorted(root.glob("case[0-9][0-9][0-9].png")):
        img = Image.open(f).convert("RGB")
        tf = root / f"{f.stem}_tumor.png"
        m = (np.array(Image.open(tf).convert("L")) > 127).astype(np.uint8) if tf.exists() else \
            np.zeros((img.height, img.width), np.uint8)
        w.add(f.stem, img, m, group=f.stem, extra={"clinical": clin.get(f.stem, {})})
    return w.finish({"source": "https://www.cancerimagingarchive.net/collection/breast-lesions-usg/",
                     "license": "CC BY 4.0", "citation": "Pawlowska et al., Scientific Data 2024",
                     "split": "seeded (0) 70/10/20 by case",
                     "notes": "tumor masks only; *_otherK.png masks (other findings) not used"})


def tn3k(raw: Path, dst: Path) -> dict:
    root = raw / "tn3k_rar/datasets/tn3k"  # TN3K.rar from HF haifan-gong/TN3K, extracted
    w = Writer(dst, ["thyroid nodule"], LESION)
    trainval = sorted((root / "trainval-image").glob("*.jpg"))
    rng = random.Random(0)
    val = set(rng.sample([f.stem for f in trainval], round(0.1 * len(trainval))))
    for part, split_fn in (("trainval", lambda s: "val" if s in val else "train"), ("test", lambda s: "test")):
        for f in sorted((root / f"{part}-image").glob("*.jpg")):
            m = (np.array(Image.open(root / f"{part}-mask" / f.name).convert("L")) > 127).astype(np.uint8)
            w.add(f"{part}_{f.stem}", Image.open(f), m, group=f"{part}_{f.stem}", split=split_fn(f.stem))
    return w.finish({"source": "HF haifan-gong/TN3K, TN3K.rar (uploaded by the TN3K author)", "license": "MIT (HF tag)",
                     "citation": "Gong et al., Thyroid region prior guided attention (TRFE-Net), CIBM 2022",
                     "split": "official test (614); val = seeded 10% of trainval",
                     "notes": "jpg masks thresholded at 127"})


def mmotu(raw: Path, dst: Path) -> dict:
    root = raw / "mmotu/x/MMOTU"          # Zenodo 21128657 (2D images + binary labels)
    lists = raw / "mmotu_lists/OTU_2d"    # official split / class lists from HF norayao/MMOTU
    names = {"train": [l.strip() for l in open(lists / "train.txt") if l.strip()],
             "test": [l.strip() for l in open(lists / "val.txt") if l.strip()]}
    cls = {}
    for f in ("train_cls.txt", "val_cls.txt"):
        for l in open(lists / f):
            p = l.split()
            if len(p) >= 2:
                cls[Path(p[0]).stem] = int(p[1])
    rng = random.Random(0)
    val = set(rng.sample(names["train"], round(0.1 * len(names["train"]))))
    w = Writer(dst, ["ovarian tumor"], LESION)
    for split, ids in names.items():
        for sid in ids:
            m = (np.array(Image.open(root / "labels" / f"{sid}.PNG")) > 0).astype(np.uint8)
            w.add(sid, Image.open(root / "images" / f"{sid}.JPG"), m, group=sid,
                  split="val" if sid in val else split, extra={"tumor_class": cls.get(sid)})
    return w.finish({"source": "images + binary labels: https://zenodo.org/records/21128657 (curated 2D MMOTU); "
                               "official train/val lists and tumour classes: HF norayao/MMOTU OTU_2d/*.txt",
                     "license": "CC BY 4.0 (Zenodo curated copy)",
                     "citation": "Zhao et al., MMOTU, arXiv:2207.06799",
                     "split": "official val (469) -> test; val = seeded 10% of official train (1000)",
                     "notes": "binary tumour masks; 8-class tumour type (0-7) kept per record; ids match the "
                              "official lists 1:1"})


def hc18(raw: Path, dst: Path) -> dict:
    from scipy import ndimage

    root = raw / "hc18/x/training_set"
    w = Writer(dst, ["fetal head"], STRUCTURE)
    bad = []
    for f in sorted(p for p in root.glob("*HC.png") if not p.name.endswith("_Annotation.png")):
        ann = np.array(Image.open(root / f"{f.stem}_Annotation.png").convert("L")) > 0
        filled = ndimage.binary_fill_holes(ann)
        if filled.sum() < 5 * ann.sum():  # outline not closed: fit the ellipse it was drawn from
            filled = fill_ellipse_contour(ann)
            if filled is None:
                bad.append(f.stem)
                continue
        fetus = re.match(r"(\d+)_", f.name).group(1)
        w.add(f.stem, Image.open(f), filled.astype(np.uint8), group=f"fetus{fetus}")
    return w.finish({"source": "https://zenodo.org/records/1327317", "license": "CC BY 4.0",
                     "citation": "van den Heuvel et al., PLoS ONE 2018 (HC18 challenge)",
                     "split": "seeded (0) 70/10/20 by fetus (NNN prefix; NNN_kHC = extra images of the same fetus); "
                              "the official test set has no public labels",
                     "notes": f"ellipse outlines filled (hole filling, or a direct least-squares ellipse fit when the "
                              f"outline is not closed); {len(bad)} dropped: {bad[:20]}"})


def psfhs(raw: Path, dst: Path) -> dict:
    import SimpleITK as sitk

    root = raw / "psfhs/x/PSFHS"
    w = Writer(dst, ["pubic symphysis", "fetal head"], STRUCTURE)
    for f in sorted((root / "image_mha").glob("*.mha")):
        img = sitk.GetArrayFromImage(sitk.ReadImage(str(f)))
        img = img[0] if img.ndim == 3 and img.shape[0] == 3 else img
        lab = sitk.GetArrayFromImage(sitk.ReadImage(str(root / "label_mha" / f.name))).astype(np.uint8)
        w.add(f.stem, _gray_rgb(img), lab, group=f.stem)
    return w.finish({"source": "https://zenodo.org/records/10969427", "license": "CC BY 4.0",
                     "citation": "Chen et al., PSFHS, Scientific Data 2024",
                     "split": "seeded (0) 70/10/20 by near-duplicate cluster (dHash <= 6); no patient ids in the source",
                     "notes": "labels 1 = pubic symphysis, 2 = fetal head (verified by area); 256x256 images"},
                    dedup=True)


def camus(raw: Path, dst: Path) -> dict:
    import zipfile

    import SimpleITK as sitk

    x = raw / "camus/x"
    if not x.exists():
        zipfile.ZipFile(raw / "camus/database_nifti.zip").extractall(x)
    w = Writer(dst, ["left ventricle", "myocardium", "left atrium"], STRUCTURE)
    for gt in sorted(x.rglob("patient*_[24]CH_E[DS]_gt.nii.gz")):
        im = gt.with_name(gt.name.replace("_gt", ""))
        a = sitk.GetArrayFromImage(sitk.ReadImage(str(im))).squeeze()
        m = sitk.GetArrayFromImage(sitk.ReadImage(str(gt))).squeeze().astype(np.uint8)
        pid = gt.name.split("_")[0]
        w.add(gt.name.replace("_gt.nii.gz", ""), _gray_rgb(a), m, group=pid)
    return w.finish({"source": "https://humanheart-project.creatis.insa-lyon.fr/database/ (CAMUS, NIfTI)",
                     "license": "CC BY-NC-SA 4.0 (non-commercial)",
                     "citation": "Leclerc et al., IEEE TMI 2019",
                     "split": "seeded (0) 70/10/20 by patient",
                     "notes": "ED and ES frames of the 2CH and 4CH views; labels 1 LV endocardium, 2 myocardium, 3 LA"})


DATASETS = {f.__name__: f for f in (busbra, busi, bus_uclm, breast_lesions, tn3k, mmotu, hc18, psfhs, camus)}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw", default="/raid/DATASETS/BioMLLMData/datasets/raw/us_bench")
    ap.add_argument("--dst", default="/raid/DATASETS/BioMLLMData/datasets/us_bench")
    ap.add_argument("--datasets", default="all")
    args = ap.parse_args()
    names = list(DATASETS) if args.datasets == "all" else args.datasets.split(",")
    summary = {}
    for n in names:
        try:
            meta = DATASETS[n](Path(args.raw), Path(args.dst) / n)
            summary[n] = {"stats": meta["stats"], "near_dup": {k: v["n"] for k, v in meta["near_duplicates_across_splits"].items()},
                          "group_leakage": len(meta["group_leakage"])}
        except Exception as e:  # keep converting the others
            summary[n] = {"error": f"{type(e).__name__}: {e}"}
        print(n, json.dumps(summary[n]), flush=True)
    write_suite(Path(args.dst))


def write_suite(dst: Path) -> dict:
    """suite.json: one entry per converted dataset (read back from every meta.json)."""
    suite = {}
    for m in sorted(dst.glob("*/meta.json")):
        meta = json.loads(m.read_text())
        suite[m.parent.name] = {"classes": meta["classes"], "boxes": meta["boxes"], "license": meta["license"],
                                "source": meta["source"], "split": meta["split"], "stats": meta["stats"],
                                "near_duplicates_across_splits": {k: v["n"] for k, v in
                                                                  meta["near_duplicates_across_splits"].items()},
                                "extra_files": [v["nodup_file"] for v in meta["near_duplicates_across_splits"].values()
                                                if "nodup_file" in v]}
    (dst / "suite.json").write_text(json.dumps(suite, indent=1))
    return suite


if __name__ == "__main__":
    main()
